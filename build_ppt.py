#!/usr/bin/env python3
"""Build Portofolio-Analisis-ISPA-Indonesia.pptx (16:9, 10 slides, ID).

Redesign design-taste: Tema T3 (oat / plum / tangerine), varian Minimalist.
Konten dan semua angka identik dengan build_ppt_legacy.py.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

# Palet T3 terkunci
GROUND = RGBColor(0xF0, 0xE8, 0xD8)   # oat
SURFACE = RGBColor(0xF8, 0xF2, 0xE4)  # permukaan terang
INK = RGBColor(0x2A, 0x1E, 0x33)      # tinta utama
PLUM = RGBColor(0x5E, 0x2A, 0x72)     # primary
TANG = RGBColor(0xEF, 0x7D, 0x1A)     # tangerine, hemat
PLUM_L = RGBColor(0x8A, 0x5C, 0xA6)   # plum muda
MUTE = RGBColor(0x6B, 0x5B, 0x74)     # teks sekunder
FONT = "Noto Sans"
CHARTS = "/home/hatch/workspace/ispa-indonesia/charts"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def textbox(slide, left, top, width, height):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx.text_frame
    tf.word_wrap = True
    return tf


def add_par(tf, text, size, color, bold=False, italic=False, first=False, align=None,
            space_after=8, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if line is not None:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def add_bullet(tf, text, size=16, first=False, color=INK):
    p = add_par(tf, text, size, color, first=first, space_after=10, line=1.12)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    pPr.set("marL", "171450")
    pPr.set("indent", "-171450")
    p.space_before = Pt(2)
    return p


def title_block(slide, text):
    """Judul konten: kotak tangerine kecil + judul ink 32pt. Tanpa eyebrow."""
    rect(slide, 0.75, 0.62, 0.16, 0.16, TANG)
    tf = textbox(slide, 1.1, 0.38, 11.6, 0.85)
    add_par(tf, text, 32, INK, bold=True, first=True, space_after=0)


def caption(slide, left, top, width, text):
    tf = textbox(slide, left, top, width, 0.4)
    add_par(tf, text, 12, MUTE, italic=True, first=True, space_after=0)


def split_slide(title, bullets, chart, cap, chart_left=False, chart_w=5.9, chart_h=3.8):
    """Layout split teks/chart, sisi bisa ditukar untuk variasi."""
    s = prs.slides.add_slide(blank)
    solid_bg(s, GROUND)
    title_block(s, title)
    if chart_left:
        bx, cx = 6.95, 0.75
    else:
        bx, cx = 0.75, 6.7
    tf = textbox(s, bx, 1.6, 5.65, 5.3)
    for i, b in enumerate(bullets):
        add_bullet(tf, b, first=(i == 0))
    s.shapes.add_picture(f"{CHARTS}/{chart}.png", Inches(cx), Inches(1.75),
                         width=Inches(chart_w), height=Inches(chart_h))
    caption(s, cx, 1.75 + chart_h + 0.18, chart_w, cap)
    return s


# ------------------------------------------------------------------ 1 Cover
# Asymmetric Split: teks kiri, angka jangkar raksasa kanan.
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
rect(s, 0, 0, 0.22, 7.5, PLUM)
tf = textbox(s, 0.85, 1.55, 6.9, 2.6)
add_par(tf, "ISPA di Indonesia 2026", 43, INK, bold=True, first=True, space_after=14, line=1.05)
add_par(tf, "Lonjakan Kasus, Sebaran Provinsi, dan Dampak Karhutla & Kualitas Udara",
        20, PLUM, line=1.2)
tf = textbox(s, 7.95, 1.7, 4.95, 2.9)
add_par(tf, "14.088.945", 50, PLUM, bold=True, first=True, space_after=6)
add_par(tf, "laporan kasus ISPA nasional", 16, INK, space_after=2)
add_par(tf, "SKDR, s.d. minggu ke-38", 14, MUTE, space_after=0)
rect(s, 7.95, 4.35, 1.5, 0.07, TANG)
tf = textbox(s, 0.85, 5.6, 11.6, 1.2)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 19, INK, bold=True, first=True, space_after=5)
add_par(tf, "Analisis Data Kesehatan Masyarakat  •  SQL + Python  •  Oktober 2026",
        15, MUTE, space_after=0)
tf = textbox(s, 0.85, 6.85, 11.6, 0.4)
add_par(tf, "Data kompilasi dari angka resmi Kemenkes / SKDR / ISPU - setiap angka bersumber.",
        12, MUTE, italic=True, first=True, space_after=0)

# ------------------------------------------------------------------ 2 Masalah
split_slide(
    "Mengapa ISPA 2026 Mendesak?",
    [
        "ISPA adalah infeksi saluran pernapasan akut: batuk, pilek, sesak napas.",
        "Laporan nasional 2026 tembus 14.088.945 laporan (SKDR, s.d. minggu ke-38).",
        "Laporan mingguan di atas 400 ribu sejak minggu ke-27 - naik ke ~420 ribu pada puncak kemarau.",
        "Bersamaan dengan karhutla dan kualitas udara yang memburuk di Sumatra dan Kalimantan.",
        "Pertanyaan: seberapa besar, di mana, siapa yang paling terdampak?",
    ],
    chart="09_mingguan_nasional",
    cap="Laporan mingguan nasional naik 5% ke puncak kemarau.",
    chart_w=5.9, chart_h=3.79,
)

# ------------------------------------------------------- 3 Data & Metode (kiri)
split_slide(
    "Data & Metode",
    [
        "Dataset KOMPILASI dari angka resmi yang dipublikasikan (Kemenkes, SKDR, ISPU, Dinkes DKI, BPS).",
        "Tidak ada file CSV mentah ISPA di data.go.id / Jakarta Open Data - setiap angka dicatat sumber + tanggalnya.",
        "Dua seri dipisah ketat: seri nasional SKDR vs seri pemantauan karhutla (cakupan dan tanggal berbeda).",
        "Pipeline: kompilasi CSV → SQLite → query SQL → chart Python → interpretasi.",
        "Perbandingan antarprovinsi dinormalisasi per 1.000 penduduk (BPS 2026).",
    ],
    chart="02_pangsa_top10",
    cap="10 provinsi teratas = 74,7% laporan nasional.",
    chart_left=True, chart_w=5.3, chart_h=3.79,
)

# ------------------------------------------------- 4 Temuan 1 (chart full-width)
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title_block(s, "Temuan 1: Kasus Sangat Terkonsentrasi")
b_left = [
    "10 provinsi teratas mencatat 10.530.385 laporan = 74,7% total nasional.",
    "4 provinsi teratas saja (Jabar, Jateng, Jakarta, Jatim) = 55,8% nasional.",
    "Jawa Barat tertinggi: 2.533.888 laporan (17,98% nasional).",
]
b_right = [
    "86,3% laporan top-10 berasal dari Pulau Jawa.",
    "Peringatan: dominasi Jawa ikut dipengaruhi jumlah penduduk dan kelengkapan pelaporan.",
]
tf = textbox(s, 0.75, 1.5, 5.9, 1.85)
for i, b in enumerate(b_left):
    add_bullet(tf, b, size=15, first=(i == 0))
tf = textbox(s, 6.95, 1.5, 5.65, 1.85)
for i, b in enumerate(b_right):
    add_bullet(tf, b, size=15, first=(i == 0))
s.shapes.add_picture(f"{CHARTS}/01_top10_absolut.png", Inches(0.75), Inches(3.35),
                     width=Inches(7.5), height=Inches(3.85))
tf = textbox(s, 8.6, 4.35, 4.0, 1.6)
add_par(tf, "74,7%", 44, PLUM, bold=True, first=True, space_after=4)
add_par(tf, "dari seluruh laporan nasional terkumpul di 10 provinsi saja", 14, MUTE,
        space_after=0, line=1.15)
caption(s, 0.75, 7.08, 7.5, "10 provinsi dengan kasus ISPA terbanyak, SKDR 2026.")

# ------------------------------------------------------- 5 Temuan 2 (chart kiri)
split_slide(
    "Temuan 2: Per Kapita, Jakarta Beban Terberat",
    [
        "Angka absolut menipu: provinsi besar otomatis mencatat kasus lebih banyak.",
        "Setelah dibagi populasi: Jakarta 150,0 kasus per 1.000 penduduk.",
        "Itu ~3x Jawa Barat (49,5) dan ~2,3x Banten (65,8).",
        "Jakarta mencatat kasus 1,16x Jawa Timur dengan penduduk hanya seperempatnya.",
        "Implikasi: surveilans dan kesiapan RS Jakarta perlu prioritas pada puncak kemarau.",
    ],
    chart="03_per_kapita",
    cap="Kasus per 1.000 penduduk (populasi BPS 2026).",
    chart_left=True, chart_w=5.9, chart_h=3.28,
)

# ------------------------------------------------ 6 Temuan 3 (band angka besar)
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title_block(s, "Temuan 3: 283.270 Kasus ISPA Karhutla")
tf = textbox(s, 0.75, 1.5, 3.6, 0.8)
add_par(tf, "283.270", 44, PLUM, bold=True, first=True, space_after=0)
tf = textbox(s, 4.35, 1.78, 7.9, 0.5)
add_par(tf, "kasus pada seri pemantauan karhutla, terpisah dari seri nasional SKDR",
        15, INK, first=True, space_after=0)
tf = textbox(s, 0.75, 2.85, 5.65, 4.2)
for i, b in enumerate([
        "Seri pemantauan karhutla (terpisah dari seri nasional): 283.270 kasus, 18 Agt - 13 Okt 2026.",
        "Tercatat di 63 kabupaten/kota, 7 provinsi - rata-rata 4.496 kasus per kab/kota.",
        "Kalsel (74.636) dan Sumsel (74.458) bersama = 52,6% kasus karhutla; selisih hanya 178 kasus.",
        "Kumulatif naik 21,1x dalam 7 minggu: 13.449 (27 Agt) → 283.270 (13 Okt).",
        "13 Oktober saja: 3.897 kasus baru harian.",
]):
    add_bullet(tf, b, first=(i == 0))
s.shapes.add_picture(f"{CHARTS}/06_karhutla_pertumbuhan.png", Inches(6.7), Inches(2.85),
                     width=Inches(5.9), height=Inches(3.69))
caption(s, 6.7, 6.72, 5.9, "Pertumbuhan kumulatif seri karhutla.")

# ------------------------------------------------------- 7 Temuan 4 (chart kiri)
split_slide(
    "Temuan 4: Hampir 1 dari 4 Kasus Karhutla adalah Balita",
    [
        "Balita (0-5 tahun): 63.711 kasus = 22,5% dari total seri karhutla.",
        "Usia di atas 5 tahun: 219.559 kasus (77,5%).",
        "Porsi balita stabil: 23,1% (21 Sep) → 22,5% (13 Okt).",
        "Anak adalah kelompok rentan utama paparan asap: saluran napas lebih kecil, frekuensi napas lebih tinggi.",
        "Respons: 2.106 tenaga kesehatan dimobilisasi; >2,8 juta masker didistribusikan.",
    ],
    chart="05_karhutla_usia",
    cap="Kasus ISPA karhutla menurut usia.",
    chart_left=True, chart_w=5.3, chart_h=3.79,
)

# ------------------------------------------------------------------ 8 Temuan 5
split_slide(
    "Temuan 5: Udara Jauh di Atas Ambang Berbahaya",
    [
        "PM2,5 di 10 wilayah terburuk: 507-1.837 (12 Okt 2026, 15:00 WIB).",
        "Barito Kuala 1.837 = 6,1 kali ambang ISPU 'berbahaya' (>300).",
        "7 dari 10 wilayah terburuk berada di Kalimantan Tengah.",
        "Kemenkes: 12 kab/kota berkategori berbahaya, 9 wilayah sangat tidak sehat.",
        "Asap karhutla mengandung PM2,5 yang masuk jauh ke saluran pernapasan (WHO).",
    ],
    chart="07_ispu",
    cap="Konsentrasi PM2,5, 10 wilayah terburuk.",
    chart_w=5.9, chart_h=3.61,
)

# --------------------------------------------------------------- 9 Rekomendasi
s = prs.slides.add_slide(blank)
solid_bg(s, GROUND)
title_block(s, "Rekomendasi")
rec_left = [
    "Lindungi balita di wilayah karhutla: masker anak, pembelajaran jarak jauh saat ISPU berbahaya, posko kesehatan.",
    "Fokus pada Kalsel dan Sumsel: 52,6% kasus karhutla - pertahankan distribusi masker dan tenaga kesehatan.",
    "Perhatikan Jakarta: beban per kapita 150,0/1.000 - perkuat surveilans puskesmas dan kesiapan RS saat puncak kemarau.",
]
rec_right = [
    "Mulai sebelum kemarau: laporan sudah >400 ribu/minggu sejak minggu ke-27.",
    "Perbaiki transparansi data: publikasikan tabel provinsi lengkap dan deret mingguan SKDR dalam format terbuka.",
]
tf = textbox(s, 0.75, 1.7, 5.9, 5.2)
for i, b in enumerate(rec_left):
    add_bullet(tf, b, size=16, first=(i == 0))
tf = textbox(s, 6.95, 1.7, 5.65, 5.2)
for i, b in enumerate(rec_right):
    add_bullet(tf, b, size=16, first=(i == 0))

# ----------------------------------------------------------------- 10 Penutup
# Designed Close: bidang plum penuh, nama besar, link GitHub, satu ajakan.
s = prs.slides.add_slide(blank)
solid_bg(s, PLUM)
rect(s, 0.85, 1.0, 1.5, 0.07, TANG)
tf = textbox(s, 0.85, 1.35, 11.6, 0.8)
add_par(tf, "Kesimpulan", 36, SURFACE, bold=True, first=True, space_after=0)
tf = textbox(s, 0.85, 2.35, 11.4, 1.9)
add_par(tf, "ISPA 2026 bukan sekadar musim batuk-pilek: 14,1 juta laporan nasional,",
        21, SURFACE, first=True, space_after=4, line=1.25)
add_par(tf, "283 ribu kasus di wilayah karhutla, dan udara 6 kali ambang berbahaya.",
        21, SURFACE, space_after=4, line=1.25)
add_par(tf, "Data yang transparan adalah langkah pertama penanganan yang tepat.",
        21, SURFACE, space_after=0, line=1.25)
tf = textbox(s, 0.85, 4.75, 11.6, 1.5)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 28, SURFACE, bold=True, first=True, space_after=6)
add_par(tf, "github.com/baghizash/ispa-indonesia-2026-analysis", 16, SURFACE, space_after=0)
tf = textbox(s, 0.85, 6.35, 11.6, 0.8)
add_par(tf, "Buka repo-nya untuk melihat analisis lengkap, query SQL, dan sumber data per angka.",
        14, RGBColor(0xDD, 0xD0, 0xEC), first=True, space_after=3)
add_par(tf, "Dataset kompilasi dari angka resmi yang dipublikasikan - lihat README untuk batasan dan caveat.",
        12, RGBColor(0xDD, 0xD0, 0xEC), italic=True, space_after=0)

prs.save("/home/hatch/workspace/ispa-indonesia/Portofolio-Analisis-ISPA-Indonesia.pptx")
prs.save("/home/hatch/workspace/your_files/Portofolio-Analisis-ISPA-Indonesia.pptx")
print("PPT saved, slides:", len(prs.slides._sldIdLst))
